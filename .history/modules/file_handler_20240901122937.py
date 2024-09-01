import os
import pandas as pd
import requests
import re
from urllib.parse import urljoin
import xml.etree.ElementTree as ET
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from tkinter import messagebox

def select_file(file_type):
    # Logic for selecting a file using tkinter's file dialog
    pass

def load_excel_urls(file_path):
    df = pd.read_excel(file_path)
    return df

def check_file_exists(file_path, file_type):
    return os.path.exists(file_path)

def replace_placeholders_with_domain(urls, base_url):
    updated_urls = []
    for url in urls:
        if re.search(r'\(\.\*\)', url):
            url = re.sub(r'\(\.\*\)', '', url)
        if not url.startswith(base_url):
            url = urljoin(base_url, url)
        updated_urls.append(url)
    return updated_urls

def parse_sitemap(sitemap_url):
    sitemap_urls = set()
    try:
        response = requests.get(sitemap_url)
        response.raise_for_status()

        if not response.content.strip():
            raise ValueError("Die Antwort der Sitemap-URL ist leer.")

        try:
            tree = ET.ElementTree(ET.fromstring(response.content))
        except ET.ParseError as e:
            raise ValueError(f"Fehler beim Parsen des XML: {str(e)}")

        root = tree.getroot()
        namespace = {'ns': 'http://www.sitemaps.org/schemas/sitemap/0.9'}

        if root.tag.endswith('sitemapindex'):
            for sitemap in root.findall('ns:sitemap/ns:loc', namespace):
                child_sitemap_url = sitemap.text.strip()
                sitemap_urls.update(parse_sitemap(child_sitemap_url))
        elif root.tag.endswith('urlset'):
            for url in root.findall('ns:url/ns:loc', namespace):
                sitemap_urls.add(url.text.strip())

    except requests.exceptions.RequestException as e:
        messagebox.showerror("Fehler beim Einlesen der Sitemap", f"HTTP Fehler: {str(e)}")
    except ValueError as e:
        messagebox.showerror("Fehler beim Einlesen der Sitemap", f"Fehler: {str(e)}")

    return sitemap_urls

def save_results_and_report(output_dir, htaccess_results, excel_results, sitemap_results, error_logs, invalid_urls):
    testergebnisse_dir = os.path.join(output_dir, 'testergebnisse')
    logs_dir = os.path.join(output_dir, 'logs')

    os.makedirs(testergebnisse_dir, exist_ok=True)
    os.makedirs(logs_dir, exist_ok=True)

    # Save results logic here, similar to the older version but modularized
    pass
